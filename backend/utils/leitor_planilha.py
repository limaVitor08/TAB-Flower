"""
Lê as planilhas da importação em lote (funcionalidade 18).
Aceita CSV e XLSX, que é o que o escritório usa hoje. O CSV é lido com o
próprio Python e o XLSX precisa do openpyxl.
Este arquivo não valida nada. Ele só transforma a planilha numa lista
com uma linha por cliente. Quem valida é o Service, com as mesmas regras
do cadastro manual.
"""
import csv
import io
import logging
logger = logging.getLogger("tabflower")
class ErroLeitura(Exception):
    """Problema no formato do arquivo, não no conteúdo das linhas."""
def _normalizar_cabecalho(nome):
    """Deixa o nome da coluna previsível.
    'Nome do Cliente ' vira 'nome_do_cliente'. Assim a planilha do
    escritório não precisa ter o cabeçalho exatamente igual ao do banco.
    """
    import unicodedata
    import re
    texto = unicodedata.normalize("NFKD", str(nome or ""))
    texto = texto.encode("ASCII", "ignore").decode("ASCII")
    texto = re.sub(r"[^A-Za-z0-9]+", "_", texto)
    return texto.strip("_").lower()
def ler_csv(conteudo_bytes):
    """Lê CSV lidando com as variações que o Excel brasileiro produz."""
    texto = None
    for codificacao in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = conteudo_bytes.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        raise ErroLeitura("Não foi possível ler o arquivo. Salve como UTF-8.")
    # O Excel em português salva separando por ponto e vírgula.
    amostra = texto[:2000]
    separador = ";" if amostra.count(";") > amostra.count(",") else ","
    leitor = csv.DictReader(io.StringIO(texto), delimiter=separador)
    if not leitor.fieldnames:
        raise ErroLeitura("A planilha está vazia ou não tem cabeçalho.")
    linhas = []
    for bruta in leitor:
        linhas.append(
            {
                _normalizar_cabecalho(k): (str(v).strip() if v is not None else "")
                for k, v in bruta.items()
                if k is not None
            }
        )
    return linhas
def ler_xlsx(conteudo_bytes):
    """Lê XLSX usando openpyxl, se estiver instalado."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ErroLeitura(
            "Leitura de XLSX exige a biblioteca openpyxl. "
            "Rode 'pip install -r requirements.txt' ou salve a planilha "
            "como CSV e envie novamente."
        )
    try:
        planilha = load_workbook(io.BytesIO(conteudo_bytes), read_only=True, data_only=True)
    except Exception as erro:
        # A mensagem na tela é generica de proposito; o motivo real do
        # openpyxl (que so ajuda quem programa) vai para o log.
        logger.error("Falha ao abrir XLSX: %s", erro)
        raise ErroLeitura("Não foi possível abrir a planilha. O arquivo pode estar corrompido.")
    aba = planilha.active
    linhas_brutas = list(aba.iter_rows(values_only=True))
    if not linhas_brutas:
        raise ErroLeitura("A planilha está vazia.")
    cabecalho = [_normalizar_cabecalho(c) for c in linhas_brutas[0]]
    linhas = []
    for bruta in linhas_brutas[1:]:
        # Ignora linhas totalmente em branco no meio da planilha.
        if all(c is None or str(c).strip() == "" for c in bruta):
            continue
        registro = {}
        for i, coluna in enumerate(cabecalho):
            if not coluna:
                continue
            valor = bruta[i] if i < len(bruta) else None
            registro[coluna] = str(valor).strip() if valor is not None else ""
        linhas.append(registro)
    return linhas
def ler(nome_arquivo, conteudo_bytes):
    """Escolhe o leitor pela extensão do arquivo."""
    nome = (nome_arquivo or "").lower()
    if nome.endswith(".csv"):
        return ler_csv(conteudo_bytes)
    if nome.endswith(".xlsx"):
        return ler_xlsx(conteudo_bytes)
    raise ErroLeitura("Envie um arquivo .csv ou .xlsx.")